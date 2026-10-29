"""Create the short committee note as an editable Google Docs-ready DOCX."""

import argparse
from pathlib import Path
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
STEM = "sanskrit_language_system_committee_note_oct7_short"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skill-dir", type=Path, required=True)
    args = parser.parse_args()
    sys.path.insert(0, str(args.skill_dir / "scripts"))
    from table_geometry import apply_table_geometry, column_widths_from_weights

    work = ROOT / "tmp/docx" / STEM
    work.mkdir(parents=True, exist_ok=True)
    output = ROOT / "output/docx" / f"{STEM}.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    raw = work / "raw.docx"
    subprocess.run([
        "pandoc", str(ROOT / "outreach/strategy" / f"{STEM}.md"),
        "-f", "markdown", "-t", "docx", "-o", str(raw),
    ], check=True)
    doc = Document(raw)

    # Google Docs preset with A4 and compact-note body rhythm (6pt, 1.10).
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(21)
        section.left_margin = section.right_margin = Mm(23)
        section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True
    for name in ("Body Text", "First Paragraph", "Compact"):
        style = doc.styles[name]
        style.base_style = normal
        style.font.name, style.font.size = "Arial", Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6 if name != "Compact" else 4)
        style.paragraph_format.line_spacing = 1.10
    for name, size, before, after in (
        ("Heading 1", 20, 20, 6), ("Heading 2", 16, 18, 6),
        ("Heading 3", 14, 16, 4),
    ):
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string("434343" if name == "Heading 3" else "000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True

    for i, paragraph in enumerate(doc.paragraphs):
        if i == 0:
            paragraph.style = normal
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.font.name, run.font.size = "Arial", Pt(26)
                run.font.bold = False
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif i == 1:
            paragraph.style = normal
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.italic = True
        elif paragraph.style.name == "Heading 2":
            paragraph.style = doc.styles["Heading 1"]
        elif paragraph.style.name == "Heading 3":
            paragraph.style = doc.styles["Heading 2"]

    for level in doc.part.numbering_part.element.xpath(".//w:lvl"):
        props = level.find(qn("w:pPr"))
        if props is None:
            props = OxmlElement("w:pPr")
            level.append(props)
        for tag in ("w:ind", "w:tabs", "w:spacing"):
            for child in props.findall(qn(tag)):
                props.remove(child)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "720")
        indent.set(qn("w:hanging"), "360")
        props.append(indent)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        props.append(tabs)
        spacing = OxmlElement("w:spacing")
        for key, value in {"before": "0", "after": "80", "line": "276", "lineRule": "auto"}.items():
            spacing.set(qn("w:" + key), value)
        props.append(spacing)

    width = (doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin) // 635
    for table in doc.tables:
        apply_table_geometry(table, column_widths_from_weights([0.25, 0.75], width), indent_dxa=0)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "insideH", "left", "right", "insideV"):
            border = OxmlElement("w:" + edge)
            border.set(qn("w:val"), "single" if edge in ("top", "bottom", "insideH") else "nil")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "DADCE0")
            borders.append(border)
        for old in table._tbl.tblPr.findall(qn("w:tblBorders")):
            table._tbl.tblPr.remove(old)
        table._tbl.tblPr.append(borders)
        for row_i, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if row_i == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.style = normal
                    paragraph.paragraph_format.space_after = Pt(4)
                    if row_i == 0:
                        for run in paragraph.runs:
                            run.bold = True

        following = table._tbl.getnext()
        if following is not None and following.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            Paragraph(following, doc).paragraph_format.space_before = Pt(6)

    # Complex-script fallback keeps Devanagari editable through Docs import.
    for run in doc.element.xpath(".//w:r"):
        props = run.get_or_add_rPr()
        fonts = props.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            props.append(fonts)
        for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        for attr in ("ascii", "hAnsi"):
            fonts.set(qn("w:" + attr), "Arial")
        fonts.set(qn("w:cs"), "Noto Sans Devanagari")
    doc.core_properties.author = "Parag Tope"
    doc.core_properties.title = "Sanskrit as a Language System"
    staged = work / "styled.docx"
    doc.save(staged)
    sanitizer = args.skill_dir / "scripts/google_docs_title_sanitize.py"
    subprocess.run([sys.executable, str(sanitizer), str(staged), "--out", str(output)], check=True)
    subprocess.run([sys.executable, str(sanitizer), str(output), "--check"], check=True)
    print(output)


if __name__ == "__main__":
    main()
